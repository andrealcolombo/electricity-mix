"""Build the electricity-mix deck (one doughnut pair per country) from monthly data.

Usage:
  python src/build_deck.py                         # latest month available, data/processed/monthly_gwh.csv
  python src/build_deck.py --end-months 8 9        # one section per period: Jan–Aug, Jan–Sep
  python src/build_deck.py --data data/sample/monthly_gwh_sample.csv
"""
import argparse
import calendar
import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Pt

from config import (BASE_YEAR, CATEGORIES, COMPARISON_YEAR, COUNTRIES, FONT, GREY,
                    LABEL_MIN_SHARE, NAVY)

SOURCE_NOTE = ("Source: ENTSO-E Transparency Platform, actual generation per production type. "
               "{period} of each year. Excludes pumped storage and other storage, and generation "
               "not metered by the transmission operator (mainly rooftop solar). "
               "Non-fossil = renewables and nuclear. Labels shown for shares of 5% or more.")

# Layout in EMU, copied from the original deck (16:9, 12192000 x 6858000)
CHART_X = {BASE_YEAR: 822960, COMPARISON_YEAR: 5120640}
CHART_Y, CHART_SIZE = 1783080, 3931920


def period_names(end_month):
    first, last = calendar.month_name[1], calendar.month_name[end_month]
    return f"{first}–{last}", f"{first[:3]}–{last[:3]}"


def add_text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """runs: list of paragraphs, each a (text, size_pt, bold, colour) tuple."""
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    for i, (text, size, bold, color) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = RGBColor.from_string(color)
    return tb


def add_doughnut(slide, year, short_period, mix):
    data = CategoryChartData()
    data.categories = list(mix.index)
    data.add_series(f"Generation {year}, {short_period} (GWh)", [round(v, 1) for v in mix.values])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Emu(CHART_X[year]), Emu(CHART_Y),
                                   Emu(CHART_SIZE), Emu(CHART_SIZE), data).chart
    chart.has_legend = False
    chart.has_title = False
    chart.plots[0].vary_by_categories = True
    # hole size is not exposed by python-pptx
    chart.plots[0]._element.find("{*}holeSize").set("val", "55")

    shares = mix / mix.sum()
    series = chart.plots[0].series[0]
    for i, (cat, share) in enumerate(shares.items()):
        pt = series.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = RGBColor.from_string(CATEGORIES[cat]["color"])
        pt.format.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        pt.format.line.width = Pt(1.25)
        if share >= LABEL_MIN_SHARE:
            tf = pt.data_label.text_frame
            tf.text = f"{share:.0%}"
            run = tf.paragraphs[0].runs[0]
            run.font.size, run.font.bold, run.font.name = Pt(9), True, FONT
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            pt.data_label.position = XL_LABEL_POSITION.CENTER


def country_slide(prs, code, end_month, mixes):
    long_period, short_period = period_names(end_month)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_text(slide, 548640, 411480, 11064240, 548640,
             [(f"{COUNTRIES[code]}: electricity generation mix, {long_period} "
               f"{BASE_YEAR} and {COMPARISON_YEAR}", 24, True, NAVY)])

    for year, mix in mixes.items():
        x = CHART_X[year]
        add_text(slide, x, 1234440, CHART_SIZE, 411480,
                 [(f"{year} ({short_period})", 18, True, NAVY)], align=PP_ALIGN.CENTER)
        add_doughnut(slide, year, short_period, mix)
        add_text(slide, x + 960119, 3364991, 2011680, 768096,
                 [(f"{mix.sum() / 1000:.1f} TWh", 22, True, NAVY), ("total generation", 11, False, GREY)],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        non_fossil = sum(v for c, v in mix.items() if CATEGORIES[c]["non_fossil"]) / mix.sum()
        add_text(slide, x, 5824728, CHART_SIZE, 320040,
                 [(f"Non-fossil: {non_fossil:.0%}", 14, True, GREY)], align=PP_ALIGN.CENTER)

    # Legend: every category present in either year, in config order
    for i, cat in enumerate(next(iter(mixes.values())).index):
        y = 2139696 + i * 384048
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(9692640), Emu(y), Emu(201168), Emu(201168))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor.from_string(CATEGORIES[cat]["color"])
        box.line.fill.background()
        add_text(slide, 10012680, y - 36576, 2011680, 292608, [(cat, 14, False, "1F2933")])

    add_text(slide, 548640, 6172200, 11064240, 307777,
             [(SOURCE_NOTE.format(period=long_period), 9, False, GREY)])


def section_slide(prs, end_month):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(slide, 548640, 2743200, 11064240, 1371600,
             [(period_names(end_month)[0].upper().replace("–", "-"), 40, True, NAVY)],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def mixes_for(df, code, end_month):
    """GWh per category for both years, Jan..end_month; drops categories that are 0 in both."""
    sub = df[(df.country == code) & (df.month <= end_month) & df.year.isin([BASE_YEAR, COMPARISON_YEAR])]
    table = (sub.pivot_table(index="category", columns="year", values="gwh", aggfunc="sum")
                .reindex(index=list(CATEGORIES), columns=[BASE_YEAR, COMPARISON_YEAR]).fillna(0))
    table = table[(table > 0).any(axis=1)]
    return {year: table[year] for year in (BASE_YEAR, COMPARISON_YEAR)}


def check_data(df, end_months):
    """Warnings printed to the console and written next to the deck."""
    warnings = []
    for code in COUNTRIES:
        for year in (BASE_YEAR, COMPARISON_YEAR):
            months = set(df[(df.country == code) & (df.year == year)].month)
            missing = [m for m in range(1, max(end_months) + 1) if m not in months]
            if missing:
                warnings.append(f"{code} {year}: no data for month(s) {missing}")
        for end in end_months:
            mixes = mixes_for(df, code, end)
            base, comp = mixes[BASE_YEAR], mixes[COMPARISON_YEAR]
            for cat in base.index:
                if base[cat] > 0 and abs(base[cat] - comp[cat]) < 0.05:
                    warnings.append(f"{code} Jan–{calendar.month_abbr[end]}: '{cat}' is identical in "
                                    f"{BASE_YEAR} and {COMPARISON_YEAR} ({base[cat]:.1f} GWh) — copy error?")
    return warnings


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default="data/processed/monthly_gwh.csv")
    ap.add_argument("--end-months", type=int, nargs="*",
                    help="last month of each period (default: latest month with data)")
    ap.add_argument("--out", default="output/electricity_mix.pptx")
    args = ap.parse_args()

    df = pd.read_csv(args.data)
    end_months = args.end_months or [int(df[df.year == COMPARISON_YEAR].month.max())]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for end in end_months:
        section_slide(prs, end)
        for code in COUNTRIES:
            country_slide(prs, code, end, mixes_for(df, code, end))

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"Wrote {out} ({len(prs.slides)} slides)")

    warnings = check_data(df, end_months)
    report = out.with_name("data_checks.md")
    report.write_text("# Data checks\n\n" + ("\n".join(f"- {w}" for w in warnings) or "No issues found.") + "\n")
    for w in warnings:
        print("WARNING:", w, file=sys.stderr)


if __name__ == "__main__":
    main()
